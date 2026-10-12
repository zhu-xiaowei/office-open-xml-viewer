from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from PIL import Image as P, ImageDraw, ImageFont
import io

def png(color, label, size=(160, 120)):
    im = P.new("RGB", size, color)
    d = ImageDraw.Draw(im)
    try:
        f = ImageFont.truetype("DejaVuSans-Bold.ttf", 22)
    except OSError:
        f = ImageFont.load_default()
    d.rectangle([0, 0, size[0]-1, size[1]-1], outline="black", width=3)
    d.text((10, size[1]//2 - 12), label, fill="white", font=f)
    path = f"{label.replace(' ', '_')}.png"
    im.save(path)
    return path

wb = Workbook()
ws = wb.active
ws.title = "Pictures"
ws["A1"] = "oneCellAnchor pictures (openpyxl default)"
ws["H1"] = "twoCellAnchor picture"
for i, (c, lbl, cell) in enumerate([("#d62728", "one-cell A", "A3"),
                                    ("#2ca02c", "one-cell B", "D3"),
                                    ("#1f77b4", "one-cell C", "A11")]):
    ws.add_image(Image(png(c, lbl)), cell)

img = Image(png("#9467bd", "two-cell"))
anchor = TwoCellAnchor()
anchor._from = AnchorMarker(col=7, row=2)
anchor.to = AnchorMarker(col=10, row=8)
img.anchor = anchor
ws.add_image(img)
wb.save("one-cell-anchor-pictures.xlsx")
